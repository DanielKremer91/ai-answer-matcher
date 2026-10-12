import io
from dataclasses import replace

import openpyxl
import pandas as pd

from aam.match import QUARTILE_LABELS
from aam.models import SOURCE, Chunk, ChunkStats, FailedUrl, SegmentedAnswer, Statement, Topic
from aam.report import (KEYWORD_COLUMN, SHEETS, answer_frame, chunks_frame, pages_frame, protocol_frame, statement_pages_frame,
                        protocol_lines, statements_frame, summary, to_excel, topics_frame)
from tests.fakes import sample_result

KEYWORD = "dachsanierung kosten"


def test_excel_has_seven_sheets_and_each_starts_with_the_keyword():
    sheets = pd.read_excel(io.BytesIO(to_excel(sample_result())), sheet_name=None)
    assert tuple(sheets) == SHEETS == ("Seiten", "Themen", "Aussagen", "Aussage x Seite", "Chunks", "Protokoll", "Erklärung")
    for name, df in sheets.items():
        assert df.columns[0] == KEYWORD_COLUMN == "Keyword/Prompt", name
        assert set(df[KEYWORD_COLUMN]) == {KEYWORD}, name


def test_pages_frame_with_sources():
    df = pages_frame(sample_result())
    assert list(df.columns) == ["Keyword/Prompt", "Rang", "URL", "Herkunft", "Seitenscore", "Bester Chunk Score",
                                "Aussagen abgedeckt", "Aussagen gesamt", "Seitenvektor"]
    assert list(df["Rang"]) == [1, 2, 3]
    assert set(df["Herkunft"]) == {"eigen", "Quelle"}
    assert set(df["Aussagen gesamt"]) == {5}
    assert all(round(v, 4) == v for v in df["Seitenscore"])


def test_topics_frame_with_sources():
    df = topics_frame(sample_result())
    assert list(df.columns) == ["Keyword/Prompt", "Thema", "Anzahl Aussagen", "Beste eigene URL",
                                "Bester eigener Chunk", "Score", "Beste Quell-URL", "Bester Quell-Chunk",
                                "Score Quelle", "Delta", "abgedeckt", "schwach", "Lücke"]
    assert list(df["Thema"]) == ["Einleitung", "Kosten im Überblick", "Förderung"]
    row = df[df["Thema"] == "Förderung"].iloc[0]
    assert (row["Anzahl Aussagen"], row["Lücke"], row["abgedeckt"]) == (2, 2, 0)
    assert row["Beste Quell-URL"] == "https://quelle.de/foerderung"


def test_statements_frame_with_sources():
    df = statements_frame(sample_result())
    assert list(df.columns) == ["Keyword/Prompt", "Thema", "Aussage (AI-Chunk)", "Wortzahl", "Beste eigene URL",
                                "Bester eigener Chunk", "Score", "Beste Quell-URL", "Bester Quell-Chunk",
                                "Score Quelle", "Delta", "Nähe zum Quell-Chunk", "Status"]
    assert len(df) == 5
    row = df[df["Aussage (AI-Chunk)"] == "Das BAFA zahlt Zuschüsse von bis zu 20 Prozent."].iloc[0]
    assert (row["Thema"], row["Wortzahl"], row["Status"]) == ("Förderung", 9, "Lücke")
    assert abs(row["Delta"] - (row["Score"] - row["Score Quelle"])) < 2e-4
    closeness = df["Nähe zum Quell-Chunk"]
    assert closeness.notna().all() and all(round(v, 4) == v for v in closeness)


def test_frames_without_sources_drop_the_source_columns():
    result = sample_result(with_sources=False)
    pages, topics, statements = pages_frame(result), topics_frame(result), statements_frame(result)
    assert "Aussagen mit bestem Chunk" in pages.columns and "Aussagen abgedeckt" not in pages.columns
    assert set(pages["Herkunft"]) == {"eigen"}
    for df in (topics, statements):
        assert not {"Beste Quell-URL", "Bester Quell-Chunk", "Score Quelle", "Delta"} & set(df.columns)
    assert "Nähe zum Quell-Chunk" not in statements.columns
    assert list(topics.columns[-4:]) == list(QUARTILE_LABELS)
    assert statements.columns[-1] == "Status"


def test_chunks_frame_lists_answer_statements_own_and_source_chunks():
    df = chunks_frame(sample_result())
    assert list(df.columns) == ["Keyword/Prompt", "Herkunft", "URL", "Chunk-Index", "Text", "Wortzahl"]
    assert df["Herkunft"].value_counts().to_dict() == {"AI-Antwort": 5, "eigen": 3, "Quelle": 2}
    assert "embedding_0" not in df.columns


def test_protocol_lines_describe_answer_chunking_and_failures():
    result = sample_result()
    failed = [FailedUrl("https://quelle.de/x", SOURCE, "HTTP 403")]
    lines = protocol_lines(result.answer, result.own_stats, result.source_stats, failed)
    assert lines[0] == ("AI-Antwort: 3 Themen, 5 Aussagen, Median 12 Wörter → Zielgröße 15 Wörter "
                        "(Segmentierung: Struktur)")
    assert lines[1].startswith("Eigener Content: 2 Seiten → 3 Chunks, Median ")
    assert "2 Absätze zusammengelegt, 0 Absätze geteilt" in lines[1]
    assert lines[2].startswith("Quellen: 1 Seite → 2 Chunks")
    assert lines[3] == "Nicht geladen (Quelle): https://quelle.de/x (HTTP 403)"


def test_protocol_lines_without_sources_and_with_prechunked_content():
    result = sample_result(with_sources=False)
    lines = protocol_lines(result.answer, result.own_stats, None, [], prechunked=True)
    assert lines[1].startswith("Eigener Content (aus Chunk-Embedding-Export): ")
    assert lines[2].startswith("Quellen: keine")


def test_protocol_frame_contains_the_run_settings():
    df = protocol_frame(sample_result())
    assert list(df.columns) == ["Keyword/Prompt", "Angabe", "Wert"]
    values = dict(zip(df["Angabe"], df["Wert"]))
    assert values["Datum"] == "2026-10-03"
    assert values["Embedding-Anbieter"] == "fake" and values["Modell"] == "fake-model"
    assert values["Dimension"] == 64 and values["Zielgröße (Wörter)"] == 15
    assert values["Segmentierung"] == "Struktur"
    assert values["Schwelle abgedeckt"] == 0.05 and values["Schwelle schwach"] == 0.15
    assert any(a == "Chunking" for a in df["Angabe"])


def test_summary_with_sources_names_top_page_and_gap_topics():
    text = summary(sample_result())
    assert text.startswith("Relevanteste eigene Seite: https://me.de/kosten (Seitenscore 0,")
    assert "von 3)" in text
    assert "Lücken: 2 Aussagen im Thema „Förderung“" in text


def test_summary_without_sources_explains_the_missing_benchmark():
    text = summary(sample_result(with_sources=False))
    assert text.startswith("Relevanteste eigene Seite: https://me.de/kosten")
    assert "keinen Benchmark" in text


def test_answer_frame_shows_every_ai_chunk():
    df = answer_frame(sample_result().answer)
    assert list(df.columns) == ["Thema", "Aussage", "Wortzahl"]
    assert len(df) == 5 and df.iloc[3]["Thema"] == "Förderung" and df.iloc[3]["Wortzahl"] == 8


NASTY_TEXTS = ["abc\x00def\x07", "=SUM(A1)", "+49 221 123", "@handle", "-5 Grad", "x" * 40000]


def _result_with_nasty_chunks():
    result = sample_result()
    own = [Chunk(url="https://me.de/nasty", origin="own", index=i, text=text, words=3)
           for i, text in enumerate(NASTY_TEXTS)]
    return replace(result, own_chunks=own)


def test_excel_export_survives_control_characters_formulas_and_long_texts():
    data = to_excel(_result_with_nasty_chunks())
    sheet = openpyxl.load_workbook(io.BytesIO(data))["Chunks"]
    header = [c.value for c in sheet[1]]
    text_col = header.index("Text") + 1
    cells = {row[text_col - 1].value: row[text_col - 1] for row in sheet.iter_rows(min_row=2)}
    values = list(cells)
    assert "abcdef" in values and not any("\x00" in v or "\x07" in v for v in values)
    for original in ("=SUM(A1)", "+49 221 123", "@handle", "-5 Grad"):
        cell = cells["'" + original]
        assert cell.data_type == "s" and cell.value == "'" + original
    long_values = [v for v in values if v.startswith("xxx")]
    assert len(long_values) == 1
    assert len(long_values[0]) <= 32767 and long_values[0].endswith("[gekürzt]")


def test_excel_cleaning_does_not_alter_the_frames_shown_in_the_ui():
    result = _result_with_nasty_chunks()
    to_excel(result)
    texts = list(chunks_frame(result)["Text"])
    for original in NASTY_TEXTS:
        assert original in texts


def _stats(documents, chunks, merged, split, empty_urls=(), median=1.0):
    return ChunkStats(documents=documents, chunks=chunks, median_words=median, p10_words=1.0, p90_words=1.0,
                      merged=merged, split=split, empty_urls=empty_urls)


def test_protocol_lines_use_singular_forms_for_counts_of_one():
    answer = SegmentedAnswer("k", (Topic(1, "T", (Statement(1, 1, "Ein Satz.", 2),)),), 1, 1.0, "structure")
    lines = protocol_lines(answer, _stats(1, 1, 1, 1, ("https://x.de",)), _stats(2, 2, 2, 2, ("a", "b")), [])
    assert lines[0] == "AI-Antwort: 1 Thema, 1 Aussage, Median 1 Wort → Zielgröße 1 Wort (Segmentierung: Struktur)"
    assert lines[1] == ("Eigener Content: 1 Seite → 1 Chunk, Median 1 Wort (p10: 1, p90: 1), "
                        "1 Absatz zusammengelegt, 1 Absatz geteilt, 1 Seite ohne Text")
    assert lines[2] == ("Quellen: 2 Seiten → 2 Chunks, Median 1 Wort (p10: 1, p90: 1), "
                        "2 Absätze zusammengelegt, 2 Absätze geteilt, 2 Seiten ohne Text")


def test_summary_keeps_aussagen_plural_for_totals_and_singular_for_single_gaps():
    text = summary(sample_result())
    assert " von 5 Aussagen abgedeckt" in text
    result = sample_result()
    topics = [replace(t, counts={**t.counts, "Lücke": 1}) if t.counts.get("Lücke") else t
              for t in result.match.topics]
    text = summary(replace(result, match=replace(result.match, topics=topics)))
    assert "1 Aussage im Thema „Förderung“" in text


def test_protocol_thresholds_are_rounded():
    result = sample_result()
    result = replace(result, meta=replace(result.meta, covered_delta=0.05000000000000001,
                                          weak_delta=0.15000000000000002))
    values = dict(zip(protocol_frame(result)["Angabe"], protocol_frame(result)["Wert"]))
    assert values["Schwelle abgedeckt"] == 0.05 and values["Schwelle schwach"] == 0.15
    assert repr(values["Schwelle schwach"]) == "0.15"


def test_protocol_line_names_the_source_target_when_it_follows_the_export():
    result = sample_result()
    lines = protocol_lines(result.answer, result.own_stats, result.source_stats, [], prechunked=True,
                           source_target_words=40)
    assert lines[2].endswith(", geschnitten mit Zielgröße 40 Wörter wie der hochgeladene Export")
    assert lines[2].startswith("Quellen: 1 Seite → 2 Chunks")
    for same in (None, 15):
        plain = protocol_lines(result.answer, result.own_stats, result.source_stats, [], source_target_words=same)
        assert "geschnitten mit" not in plain[2]
    without_sources = protocol_lines(result.answer, result.own_stats, None, [], source_target_words=40)
    assert "geschnitten mit" not in without_sources[2]


def test_protocol_frame_shows_the_source_target():
    result = sample_result()
    meta = replace(result.meta, source_target_words=40, prechunked=True)
    lines = list(protocol_frame(replace(result, meta=meta))["Wert"])
    assert any("geschnitten mit Zielgröße 40 Wörter wie der hochgeladene Export" in str(line) for line in lines)
    same = replace(result, meta=replace(result.meta, source_target_words=0))
    assert not any("geschnitten mit" in str(line) for line in protocol_frame(same)["Wert"])


def test_prechunked_protocol_line_has_no_merge_and_split_counts():
    result = sample_result()
    lines = protocol_lines(result.answer, _stats(3, 7, 2, 1, median=12.4), None, [], prechunked=True)
    assert lines[1] == ("Eigener Content (aus Chunk-Embedding-Export): 3 Seiten → 7 Chunks, "
                        "Median 12 Wörter (p10: 1, p90: 1)")
    single = protocol_lines(result.answer, _stats(1, 1, 0, 0, median=1.0), None, [], prechunked=True)
    assert single[1] == ("Eigener Content (aus Chunk-Embedding-Export): 1 Seite → 1 Chunk, "
                         "Median 1 Wort (p10: 1, p90: 1)")


def test_pages_without_text_are_listed_per_origin_and_capped():
    result = sample_result()
    own = _stats(5, 0, 0, 0, tuple(f"https://me.de/leer{i}" for i in range(22)))
    source = _stats(2, 1, 0, 0, ("https://quelle.de/leer",))
    failed = [FailedUrl("https://quelle.de/x", SOURCE, "HTTP 403")]
    lines = protocol_lines(result.answer, own, source, failed)
    empties = lines[3:-1]  # between the stats lines and the failed URLs
    assert len(empties) == 22
    assert empties[0] == "Ohne Text (eigen): https://me.de/leer0"
    assert empties[19] == "Ohne Text (eigen): https://me.de/leer19"
    assert empties[20] == "… und 2 weitere"
    assert empties[21] == "Ohne Text (Quelle): https://quelle.de/leer"
    assert lines[-1] == "Nicht geladen (Quelle): https://quelle.de/x (HTTP 403)"


def test_pages_without_text_reach_the_excel_protocol_sheet():
    result = sample_result()
    result = replace(result, own_stats=_stats(3, 3, 0, 0, ("https://me.de/leer",)))
    assert "Ohne Text (eigen): https://me.de/leer" in list(protocol_frame(result)["Wert"])
    sheet = pd.read_excel(io.BytesIO(to_excel(result)), sheet_name="Protokoll")
    assert "Ohne Text (eigen): https://me.de/leer" in list(sheet["Wert"])


def test_summary_hints_at_the_minimum_when_no_status_can_be_given():
    from aam.pipeline import analyze, prepare
    from tests.fakes import OWN_DOCS, FakeProvider

    short = "Eine Dachsanierung kostet viel.\nSie dauert mehrere Wochen.\nDie Planung braucht Zeit."
    prepared = prepare("dach", short, own_docs=OWN_DOCS)
    result = analyze(prepared, FakeProvider(), None)
    assert len(result.match.statements) == 3 and not result.match.has_sources
    assert summary(result).endswith(" Für einen Status sind mindestens 4 Aussagen oder Quell-URLs nötig.")
    assert "mindestens 4 Aussagen" not in summary(sample_result(with_sources=False))  # 5 statements
    assert "mindestens 4 Aussagen" not in summary(sample_result())


def test_statement_pages_frame_lists_each_statement_against_each_page():
    from aam.report import statement_pages_frame

    df = statement_pages_frame(sample_result())
    assert list(df.columns) == ["Keyword/Prompt", "Thema", "Aussage (AI-Chunk)", "Rang", "URL", "Herkunft",
                                "Bester Chunk Score", "Seiten-Relevanz", "Bester Chunk"]
    assert len(df) == 5 * 3  # five statements, two own pages and one source
    rows = df[df["Aussage (AI-Chunk)"] == "Die KfW fördert die Dachdämmung mit zinsgünstigen Krediten."]
    assert list(rows["Rang"]) == [1, 2, 3]
    top = rows.iloc[0]
    assert (top["URL"], top["Herkunft"], top["Thema"]) == ("https://quelle.de/foerderung", "Quelle", "Förderung")
    assert top["Bester Chunk"] == "Die KfW fördert die Dachdämmung mit zinsgünstigen Krediten."
    assert all(round(v, 4) == v for v in df["Seiten-Relevanz"])


def test_statement_pages_matrix_has_one_row_per_statement_and_one_column_per_page():
    from aam.report import statement_pages_matrix

    matrix = statement_pages_matrix(sample_result())
    ranked = [p.url + (" (Quelle)" if p.origin == SOURCE else "") for p in sample_result().match.pages]
    assert list(matrix.columns) == ["Thema", "Aussage"] + ranked  # pages in the order of the page ranking
    assert set(ranked) == {"https://me.de/kosten", "https://me.de/blog", "https://quelle.de/foerderung (Quelle)"}
    assert len(matrix) == 5
    row = matrix[matrix["Aussage"] == "Die KfW fördert die Dachdämmung mit zinsgünstigen Krediten."].iloc[0]
    assert row["https://quelle.de/foerderung (Quelle)"] == 1.0
    assert row["https://me.de/kosten"] < 0.85


def test_statement_pages_matrix_is_skipped_for_many_pages():
    from aam.report import statement_pages_matrix

    assert statement_pages_matrix(sample_result(), max_pages=2) is None


def test_protocol_line_names_switched_off_chunking_steps():
    from aam.chunk import chunk_documents
    from aam.models import OWN, Document

    _, stats = chunk_documents([Document("u", "Ein kurzer Absatz.\nNoch ein kurzer Absatz.", OWN)], 20,
                               merge_short=False, split_long=False)
    result = sample_result()
    line = protocol_lines(result.answer, stats, None, [])[1]
    assert line.endswith("kurze Absätze nicht zusammengelegt, lange Absätze nicht geteilt")


def test_protocol_frame_reports_how_page_vectors_were_built():
    df = protocol_frame(sample_result())
    values = dict(zip(df["Angabe"], df["Wert"]))
    assert values["Seitenvektoren"] == "3 Seiten komplett eingebettet, 0 als Durchschnitt ihrer Chunks"


def test_protocol_line_shows_the_content_factor():
    from aam.segment import segment_answer
    from tests.fakes import ANSWER
    result = sample_result()
    answer = segment_answer("dachsanierung kosten", ANSWER, content_factor=1.5)
    line = protocol_lines(answer, result.own_stats, result.source_stats, [])[0]
    assert "Median 12 Wörter × 1,5 → Zielgröße 18 Wörter" in line


def test_pages_frame_says_how_each_page_vector_was_built():
    result = sample_result()
    df = pages_frame(result)
    assert list(df.columns)[-1] == "Seitenvektor"
    assert set(df["Seitenvektor"]) == {"komplett"}
    import dataclasses
    averaged = dataclasses.replace(result, page_vectors={})
    assert set(pages_frame(averaged)["Seitenvektor"]) == {"Durchschnitt der Chunks"}


def test_overlong_pages_lists_pages_above_the_model_limit():
    from aam.report import overlong_pages

    texts = {("own", "https://me.de/kurz"): "kurz", ("own", "https://me.de/lang"): " ".join(["wort"] * 1300)}
    assert overlong_pages(texts, "gemini") == [
        "https://me.de/lang: 1300 Wörter, über der Grenze von 1200 Wörtern für gemini → Seitenvektor als Durchschnitt der Chunks"]
    assert overlong_pages(texts, "openai") == []


def test_explanation_sheet_covers_every_column_of_every_sheet():
    from aam.report import explanation_frame

    for result in (sample_result(), sample_result(with_sources=False)):
        legend = explanation_frame(result)
        assert list(legend.columns) == ["Keyword/Prompt", "Blatt", "Spalte", "Bedeutung"]
        explained = {(sheet, column) for sheet, column in zip(legend["Blatt"], legend["Spalte"])}
        frames = {"Seiten": pages_frame(result), "Themen": topics_frame(result), "Aussagen": statements_frame(result),
                  "Aussage x Seite": statement_pages_frame(result), "Chunks": chunks_frame(result),
                  "Protokoll": protocol_frame(result)}
        for sheet, frame in frames.items():
            for column in frame.columns:
                assert (sheet, column) in explained or (sheet, column) == (sheet, "Keyword/Prompt"), (sheet, column)
        assert ("alle", "Keyword/Prompt") in explained
        assert all(len(text) > 15 for text in legend["Bedeutung"])
